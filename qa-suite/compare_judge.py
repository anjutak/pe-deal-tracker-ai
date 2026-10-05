"""
Compare your hand labels (judge_labels.xlsx) with the Claude judge's verdicts.

Usage:
    python compare_judge.py
    python compare_judge.py judge_labels.xlsx results/judge_results_run5.json

Prints agreement, a confusion table and every disagreement, and saves
results/judge_agreement.json.
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

LABELS_FILE = Path(sys.argv[1] if len(sys.argv) > 1 else "judge_labels.xlsx")
JUDGE_FILE = Path(sys.argv[2] if len(sys.argv) > 2 else "results/judge_results_run5.json")


def load_human_labels(path: Path) -> dict:
    ws = load_workbook(path, data_only=True)["Labels"]
    labels = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        case_id, verdict, notes = row[0], row[5], row[6]
        if not case_id:
            break
        if verdict:
            labels[case_id] = {"verdict": str(verdict).strip().upper(), "notes": notes or ""}
    return labels


def cohens_kappa(pairs: list[tuple[bool, bool]]) -> float:
    """Agreement corrected for chance (1.0 = perfect, 0 = no better than chance)."""
    n = len(pairs)
    observed = sum(h == j for h, j in pairs) / n
    h_pass = sum(h for h, _ in pairs) / n
    j_pass = sum(j for _, j in pairs) / n
    expected = h_pass * j_pass + (1 - h_pass) * (1 - j_pass)
    return 1.0 if expected == 1 else (observed - expected) / (1 - expected)


def main():
    human = load_human_labels(LABELS_FILE)
    judge = {r["id"]: r for r in json.loads(JUDGE_FILE.read_text(encoding="utf-8"))["results"]}

    if not human:
        sys.exit("No verdicts found - fill in the 'Your verdict' column in judge_labels.xlsx first.")

    rows, pairs = [], []
    for case_id, h in human.items():
        if case_id not in judge:
            print(f"Warning: {case_id} not in {JUDGE_FILE}, skipped")
            continue
        j = judge[case_id]
        human_pass = h["verdict"] == "PASS"
        pairs.append((human_pass, j["passed"]))
        rows.append({
            "id": case_id,
            "human": "PASS" if human_pass else "FAIL",
            "judge": "PASS" if j["passed"] else "FAIL",
            "agree": human_pass == j["passed"],
            "judge_scores": f"C{j['correctness']} G{j['grounding']} S{j['scope']}",
            "judge_issues": j["issues"],
            "human_notes": h["notes"],
        })

    n = len(pairs)
    agree = sum(r["agree"] for r in rows)
    both_fail = sum(1 for h, j in pairs if not h and not j)
    both_pass = sum(1 for h, j in pairs if h and j)
    judge_missed = sum(1 for h, j in pairs if not h and j)      # you failed it, judge passed it
    judge_too_strict = sum(1 for h, j in pairs if h and not j)  # you passed it, judge failed it
    kappa = cohens_kappa(pairs)

    print(f"\nJudge vs. human on {n} labeled answers")
    print("=" * 50)
    print(f"Agreement:        {agree}/{n} ({agree / n:.0%})")
    print(f"Cohen's kappa:    {kappa:.2f}   (above 0.8 = strong agreement)")
    print()
    print("                    Judge PASS   Judge FAIL")
    print(f"  You said PASS        {both_pass:>3}          {judge_too_strict:>3}   <- judge too strict")
    print(f"  You said FAIL        {judge_missed:>3}          {both_fail:>3}")
    print("                        ^ judge missed a bug")

    disagreements = [r for r in rows if not r["agree"]]
    if disagreements:
        print("\nDisagreements (look at each one - who was right?)")
        for r in disagreements:
            print(f"\n  {r['id']}: you={r['human']}  judge={r['judge']} ({r['judge_scores']})")
            print(f"    Judge issues: {r['judge_issues'] or 'none'}")
            print(f"    Your notes:   {r['human_notes'] or '-'}")
    else:
        print("\nNo disagreements.")

    out = Path("results/judge_agreement.json")
    out.parent.mkdir(exist_ok=True)
    out.write_text(json.dumps({
        "labeled": n, "agreement": agree, "agreement_rate": round(agree / n, 3),
        "cohens_kappa": round(kappa, 3), "judge_missed_bug": judge_missed,
        "judge_too_strict": judge_too_strict, "rows": rows,
    }, indent=2), encoding="utf-8")
    print(f"\nSaved to {out}")


if __name__ == "__main__":
    main()
